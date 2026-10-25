"""Gera a planilha de composição do time de IA do Bytebank.

Saídas (na mesma pasta deste script):
- time_de_ia.csv      -> aba 1: composição do time
- carreira_em_y.csv   -> aba 2: caminho de gestão x caminho de especialista
- time_de_ia.xlsx     -> as duas abas num único arquivo (opcional, requer openpyxl)

Os CSVs são a fonte de verdade versionada; o .xlsx é só um atalho para importar
as duas abas de uma vez no Google Sheets.
"""

import csv
from pathlib import Path

PASTA = Path(__file__).parent

COLUNAS_TIME = [
    "Cargo",
    "Responsabilidades principais",
    "Nível esperado (Jr/Pl/Sr)",
    "Relação com o projeto de agentes IA",
    "Por que é necessário no Bytebank",
]

TIME = [
    [
        "Engenheiro(a) de Dados",
        "Construir e manter pipelines de ingestão (ETL/ELT) das políticas, FAQ e regras de cartão; garantir qualidade, versionamento e linhagem dos dados; automatizar a atualização da base de conhecimento.",
        "Pl",
        "Ingere o politicas_bytebank.csv, valida o esquema (id, dominio, secao, nivel_acesso, conteudo) e mantém a rotina que recria o vector store (ChromaDB) quando uma política muda.",
        "Um RAG só é tão confiável quanto a base que consulta: política desatualizada no índice vira resposta errada ao cliente, com risco regulatório. A linhagem dos dados torna cada resposta auditável.",
    ],
    [
        "Analista de Dados",
        "Definir e acompanhar KPIs do atendimento (taxa de resolução, transbordo para humano, satisfação); construir dashboards; analisar as perguntas mais frequentes dos clientes.",
        "Jr",
        "Mede se o chatbot resolve de fato, aponta perguntas sem resposta na base (lacunas de conteúdo) e ajuda a montar o dataset de validação com perguntas reais.",
        "A diretoria precisa de evidência de valor e de onde o chatbot falha. Sem métricas, o projeto não se sustenta nem é priorizado.",
    ],
    [
        "Cientista de Dados",
        "Desenhar experimentos e métricas de avaliação; comparar estratégias (chunking, modelo de embeddings, re-ranking); analisar vieses e qualidade estatística dos resultados.",
        "Pl",
        "Conduz a avaliação com e sem RAG, define o protocolo de LLM como juiz, mede a qualidade da recuperação e investiga respostas enviesadas entre perfis de clientes.",
        "Decisões técnicas num banco precisam de evidência mensurável e reprodutível, não de intuição; também é quem detecta tratamento desigual entre clientes (imparcialidade).",
    ],
    [
        "Engenheiro(a) de IA",
        "Projetar e implementar a arquitetura RAG e multiagente; integrar LLMs (LangChain/LangGraph); escrever prompts e guardrails; implementar protocolos A2A e MCP e os pontos de Human-in-the-Loop.",
        "Sr",
        "Constrói o núcleo do sistema: retriever com filtro por nivel_acesso, re-ranking, supervisor/roteador, agentes especializados e servidor MCP.",
        "É o papel que transforma a arquitetura em produto. Pela complexidade (vários agentes, protocolos, dados sensíveis) e pelo risco de erro com clientes, exige senioridade.",
    ],
    [
        "DPO (Encarregado(a) de Proteção de Dados)",
        "Garantir conformidade com a LGPD; elaborar o Relatório de Impacto à Proteção de Dados (RIPD); ser canal com a ANPD e com os titulares; aprovar finalidades, retenção e compartilhamento de dados pessoais.",
        "Sr",
        "Aprova quais dados podem entrar na base do RAG e nos prompts, valida o envio de dados a provedores externos de LLM e define as regras de nível de acesso (público, interno, restrito).",
        "A LGPD (art. 41) exige um encarregado; um banco trata dados pessoais e financeiros, e um vazamento via chatbot geraria multa, dano reputacional e quebra de confiança.",
    ],
    [
        "Engenheiro(a) de MLOps/LLMOps",
        "Automatizar deploy de modelos e prompts (CI/CD); versionar prompts, modelos e índices; monitorar custo, latência e qualidade; observabilidade (logs, traces, snapshots) e rollback.",
        "Pl",
        "Implementa os quatro pilares de LLM Ops (gerenciamento, otimização, automação, qualidade) e mantém a esteira que reavalia o RAG a cada mudança.",
        "Sustentável e auditável significa reproduzir qualquer resposta passada e detectar degradação antes do cliente. Isso exige automação e monitoramento contínuos.",
    ],
    [
        "Product Manager de IA",
        "Priorizar casos de uso com base em valor e risco; conectar negócio, jurídico/compliance e tecnologia; definir critérios de sucesso e regras de transbordo para atendimento humano.",
        "Sr",
        "Define o escopo de cada agente (conta corrente, cartão, suporte), quais ações exigem aprovação humana (ex.: cartão Platinum, contestação de fraude) e o roadmap.",
        "Evita construir tecnologia sem problema de negócio e garante que limites de automação sejam decididos com o negócio e o compliance, não só pela engenharia.",
    ],
]

COLUNAS_CARREIRA = [
    "Cargo",
    "Caminho que tende a seguir",
    "Por quê",
    "Evolução no caminho de especialista",
    "Evolução no caminho de gestão",
]

CARREIRA = [
    [
        "Engenheiro(a) de Dados",
        "Especialista",
        "O valor está na profundidade técnica (arquitetura de dados, performance, qualidade) que fica mais rara com a senioridade.",
        "Sr → Staff Data Engineer → Arquiteto(a) de Dados",
        "Sr → Tech Lead de Dados → Gerente de Engenharia de Dados",
    ],
    [
        "Analista de Dados",
        "Gestão",
        "Trabalha perto das áreas de negócio e da diretoria, traduzindo métricas em decisões; essa ponte favorece a liderança de times de BI.",
        "Pl → Sr → Especialista em Analytics",
        "Sr → Coordenador(a) de BI/Analytics → Gerente de Dados",
    ],
    [
        "Cientista de Dados",
        "Especialista",
        "Pesquisa, experimentação e rigor estatístico são competências que se aprofundam; o Bytebank precisa de referência técnica em avaliação de modelos.",
        "Sr → Cientista Principal → Referência em IA Responsável",
        "Sr → Líder de Ciência de Dados → Head de Dados",
    ],
    [
        "Engenheiro(a) de IA",
        "Especialista",
        "A área muda rápido (RAG, agentes, protocolos A2A/MCP); o banco ganha mais com alguém que define padrões técnicos do que com mais um gestor.",
        "Sr → Staff AI Engineer → Arquiteto(a) de IA",
        "Sr → Tech Lead de IA → Gerente de Engenharia de IA",
    ],
    [
        "DPO (Encarregado(a) de Proteção de Dados)",
        "Gestão",
        "Responde institucionalmente pela conformidade, coordena áreas (jurídico, segurança, TI) e se relaciona com a ANPD; é um papel de liderança e governança.",
        "Sr → Especialista em Privacidade e Proteção de Dados",
        "DPO → Head de Privacidade → Diretor(a) de Compliance",
    ],
    [
        "Engenheiro(a) de MLOps/LLMOps",
        "Especialista",
        "Plataforma, automação e observabilidade são competências técnicas profundas e transversais a todos os times de IA.",
        "Sr → Staff Platform Engineer → Arquiteto(a) de Plataforma de IA",
        "Sr → Tech Lead de Plataforma → Gerente de Plataforma",
    ],
    [
        "Product Manager de IA",
        "Gestão",
        "Já atua coordenando pessoas e prioridades sem autoridade formal; a evolução natural é liderar outros PMs e o portfólio de produtos de IA.",
        "Sr → Principal PM (referência em produto de IA)",
        "Sr → Group PM → Head de Produto de IA",
    ],
]


def salvar_csv(nome, colunas, linhas):
    with open(PASTA / nome, "w", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(colunas)
        escritor.writerows(linhas)


def salvar_xlsx(nome):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    abas = [("Time de IA", COLUNAS_TIME, TIME), ("Carreira em Y", COLUNAS_CARREIRA, CARREIRA)]
    for i, (titulo, colunas, linhas) in enumerate(abas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = titulo
        ws.append(colunas)
        for linha in linhas:
            ws.append(linha)
        for celula in ws[1]:
            celula.font = Font(bold=True)
        for coluna in ws.columns:
            ws.column_dimensions[coluna[0].column_letter].width = 45
            for celula in coluna:
                celula.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    wb.save(PASTA / nome)


if __name__ == "__main__":
    salvar_csv("time_de_ia.csv", COLUNAS_TIME, TIME)
    salvar_csv("carreira_em_y.csv", COLUNAS_CARREIRA, CARREIRA)
    try:
        salvar_xlsx("time_de_ia.xlsx")
    except ImportError:
        print("openpyxl não instalado: .xlsx não gerado (os CSVs foram gerados).")
    print("Planilhas geradas em", PASTA)
